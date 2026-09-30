import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Clean, plain, readable colors
    BG_DARK = RGBColor(15, 23, 42)       # Slate 900
    CARD_BG = RGBColor(30, 41, 59)       # Slate 800
    BORDER_CLR = RGBColor(51, 65, 85)    # Slate 700
    TEXT_MAIN = RGBColor(255, 255, 255)  # White
    TEXT_MUTED = RGBColor(203, 213, 225) # Slate 300
    ACCENT_BLUE = RGBColor(56, 189, 248) # Sky blue
    ACCENT_GREEN = RGBColor(74, 222, 128)# Light green
    ACCENT_AMBER = RGBColor(251, 191, 36)# Warm yellow

    blank_layout = prs.slide_layouts[6]

    def set_slide_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_DARK
        bg.line.fill.background()
        return bg

    def add_header(slide, step_label, title_text, subtitle_text=""):
        tag = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.5), Inches(2.2), Inches(0.35))
        tag.fill.solid()
        tag.fill.fore_color.rgb = CARD_BG
        tag.line.color.rgb = BORDER_CLR
        tag.line.width = Pt(1)
        tf_tag = tag.text_frame
        p_tag = tf_tag.paragraphs[0]
        p_tag.text = step_label.upper()
        p_tag.font.size = Pt(9)
        p_tag.font.bold = True
        p_tag.font.color.rgb = ACCENT_BLUE
        p_tag.alignment = PP_ALIGN.CENTER

        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.9), Inches(11.7), Inches(0.8))
        tf = title_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.size = Pt(22)
        p.font.bold = True
        p.font.color.rgb = TEXT_MAIN

        if subtitle_text:
            p_sub = tf.add_paragraph()
            p_sub.text = subtitle_text
            p_sub.font.size = Pt(11)
            p_sub.font.color.rgb = TEXT_MUTED

    def add_card(slide, left, top, width, height, title, items, badge="", accent=ACCENT_BLUE):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = BORDER_CLR
        card.line.width = Pt(1.2)

        tb = slide.shapes.add_textbox(Inches(left + 0.25), Inches(top + 0.2), Inches(width - 0.5), Inches(height - 0.4))
        tf = tb.text_frame
        tf.word_wrap = True

        p0 = tf.paragraphs[0]
        p0.text = f"[{badge}]  {title}" if badge else title
        p0.font.size = Pt(14)
        p0.font.bold = True
        p0.font.color.rgb = accent

        for item in items:
            p = tf.add_paragraph()
            p.text = f"• {item}"
            p.font.size = Pt(11)
            p.font.color.rgb = TEXT_MUTED
            p.space_before = Pt(6)

    # -------------------------------------------------------------
    # SLIDE 1: Title Slide (Simple & Plain)
    # -------------------------------------------------------------
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1)

    decor = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(11.733), Inches(4.5))
    decor.fill.solid()
    decor.fill.fore_color.rgb = CARD_BG
    decor.line.color.rgb = BORDER_CLR

    tb = s1.shapes.add_textbox(Inches(1.2), Inches(2.0), Inches(11.0), Inches(3.6))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "vvEntra"
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = TEXT_MAIN

    p2 = tf.add_paragraph()
    p2.text = "A Safe Marketplace to Buy & Sell Software Businesses"
    p2.font.size = Pt(18)
    p2.font.color.rgb = ACCENT_BLUE
    p2.space_before = Pt(6)

    p3 = tf.add_paragraph()
    p3.text = "A simple, fast website where developers and investors trade apps and software safely without getting scammed."
    p3.font.size = Pt(13)
    p3.font.color.rgb = TEXT_MUTED
    p3.space_before = Pt(14)

    p4 = tf.add_paragraph()
    p4.text = "Live Website: https://rohithagarwal.github.io/vventra/  |  Repository: github.com/rohithagarwal/vventra"
    p4.font.size = Pt(11)
    p4.font.color.rgb = ACCENT_GREEN
    p4.space_before = Pt(24)

    s1.notes_slide.notes_text_frame.text = (
        "Hello everyone. Today I'm presenting vvEntra. "
        "In simple terms, vvEntra is an online marketplace where people can buy and sell software businesses safely. "
        "I rebuilt the website using simple, clean code with zero extra libraries, made it load fast, "
        "and added a safe 4-step handover system so neither buyers nor sellers get scammed."
    )

    # -------------------------------------------------------------
    # SLIDE 2: Core Idea, Problem & Benefits
    # -------------------------------------------------------------
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2)
    add_header(s2, "1. Core Idea", "What is vvEntra, What It Solves & Why It Helps", "A simple explanation of the business and who it helps.")

    add_card(s2, 0.8, 1.9, 3.6, 5.0, "The Problem It Solves", [
        "Expensive Middlemen: Traditional brokers charge huge fees (10% to 15%) and take 6 to 12 months.",
        "Fake Numbers: Public websites are full of fake revenue screenshots and stolen code.",
        "Fear of Getting Cheated: Buyers are afraid to pay first and get ghosted. Sellers are afraid to give away code without getting paid."
    ], badge="THE PROBLEM", accent=ACCENT_AMBER)

    add_card(s2, 4.86, 1.9, 3.6, 5.0, "Benefits for the User", [
        "For Buyers: Real earnings checked with Stripe (no fake screenshots); code is checked for bugs; money stays safe in escrow.",
        "For Sellers: Much lower fees (3% to 7% instead of 15%, saving $50k+); code is never handed over until the buyer's money is locked.",
        "Fast Deals: Deals close in about 18 days instead of almost a year."
    ], badge="FOR USERS", accent=ACCENT_GREEN)

    add_card(s2, 8.93, 1.9, 3.6, 5.0, "Benefits for Us (Our Business)", [
        "Simple Business Model: We take a small 3% to 7% fee on every sale (e.g. $30,000 profit on a $500,000 deal).",
        "Extra Services: Revenue from code audits, escrow fees, and private deal room access.",
        "Valuable Market Data: We build a real database of how much software companies actually sell for."
    ], badge="FOR OUR PLATFORM", accent=ACCENT_BLUE)

    s2.notes_slide.notes_text_frame.text = (
        "Point 1: What is the core idea? "
        "Think of it like a trusted real estate agent, but for apps and software instead of houses. "
        "It solves three simple problems: high broker fees, fake earnings screenshots, and the fear of getting scammed. "
        "Buyers get real, verified numbers and safe escrow. Sellers save tens of thousands in fees and protect their code. "
        "And for our platform, we make a clear 3% to 7% fee on every successful sale."
    )

    # -------------------------------------------------------------
    # SLIDE 3: Major Functionalities (Plain & Simple)
    # -------------------------------------------------------------
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3)
    add_header(s3, "2. Main Features", "Each Major Functionality of the Website", "What users can actually do on the website.")

    add_card(s3, 0.8, 1.9, 3.6, 5.0, "Market Tools & Dashboard", [
        "Buyer / Seller Switch: One click changes the view to show what buyers care about (returns) or what sellers care about (how much their app is worth).",
        "4-Box Market Grid: Groups apps into 4 simple buckets (Big apps, Fast growers, Niche tools, Basic utilities). Clicking any box filters the apps.",
        "Supply & Demand Graph: Shows which software categories have the most buyers waiting."
    ], badge="DASHBOARD", accent=ACCENT_BLUE)

    add_card(s3, 4.86, 1.9, 3.6, 5.0, "Valuation & Marketplace", [
        "Price Calculator: A tool for sellers. You type in your revenue, and it instantly tells you what your app is worth and your chances of selling.",
        "Browse Marketplace: Search and filter apps by category (AI, Finance, Health) with real monthly earnings, profit, and programming languages shown clearly."
    ], badge="BUY & SELL", accent=ACCENT_GREEN)

    add_card(s3, 8.93, 1.9, 3.6, 5.0, "Private Room & Calculator", [
        "Private Diligence Room: A locked room where buyers can view private tax files and code stats after signing an agreement.",
        "Code Walkthrough Video: A video screen where the creator explains how the code works.",
        "Fee Calculator: A slider showing exactly how many thousands of dollars you save compared to traditional brokers."
    ], badge="DEAL ROOM", accent=ACCENT_AMBER)

    s3.notes_slide.notes_text_frame.text = (
        "Point 2: Here are the main features of the website. "
        "First, a top switch that flips between Buyer and Seller mode. "
        "Second, a dashboard with a 4-box grid where users can click to filter companies, and charts showing market demand. "
        "Third, a calculator that tells founders what their software is worth. "
        "Fourth, a clean marketplace to search apps. "
        "And fifth, a private data room to view sensitive business records safely."
    )

    # -------------------------------------------------------------
    # SLIDE 4: Tech Stack & Adding AI Agents
    # -------------------------------------------------------------
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4)
    add_header(s4, "3. Technology", "The Tech Stack & How AI Agents Can Be Added", "Simple code that makes the site fast and easy for AI to work with.")

    add_card(s4, 0.8, 1.9, 5.7, 5.0, "The Tech Stack Used", [
        "Pure HTML5 & CSS3: Clean layout and styling using modern CSS variables. No heavy frameworks like Tailwind or Bootstrap.",
        "Vanilla JavaScript: Pure code for page navigation, search, and charts. Zero npm packages, zero React, and zero build steps.",
        "GitHub Pages Hosting: Free global hosting directly from GitHub. Fast loading in under 120ms with zero server crashes."
    ], badge="SIMPLE TECH STACK", accent=ACCENT_BLUE)

    add_card(s4, 6.8, 1.9, 5.7, 5.0, "How It Helps Us Add AI Agents", [
        "Single Data File (data.js): All listings and numbers live in one clean file. AI agents can easily read and add new listings without breaking UI code.",
        "No Hidden React Code: Because the HTML is clean and direct, AI browser tools can easily click buttons, search listings, and read text.",
        "AI Agents We Can Add Easily:",
        "  1. Code Audit AI: Automatically scans GitHub repos and gives a code safety score.",
        "  2. Price AI: Checks recent sales and updates valuations automatically.",
        "  3. Q&A AI Chatbot: Sits inside the private deal room to answer buyers' questions about the software.",
        "  4. Handover AI: Checks that domains and passwords are moved correctly before money is sent."
    ], badge="ADDING AI AGENTS", accent=ACCENT_GREEN)

    s4.notes_slide.notes_text_frame.text = (
        "Point 3: For our tech stack, we kept it 100% simple: pure HTML, CSS, and Vanilla JavaScript with zero extra libraries. "
        "Why is this helpful for adding AI agents? "
        "Because all data is stored in one clean file (data.js), and the web code is completely clean without complex React layers. "
        "AI agents can easily connect to it: an AI can scan GitHub code, an AI can automatically update prices, "
        "and an AI assistant can answer buyers' questions inside the private data room."
    )

    # -------------------------------------------------------------
    # SLIDE 5: Innovation Compared to Original Website
    # -------------------------------------------------------------
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5)
    add_header(s5, "4. Innovation", "Innovation Added (Compared to Original Website)", "How our site is safer and more functional than the reference site.")

    add_card(s5, 0.8, 1.9, 5.7, 5.0, "The Original Website (vventrabyaj)", [
        "Flashy Graphics Over Substance: Used heavy neon purple colors and glows that looked like a video game or crypto token.",
        "Buttons Did Not Work: Many cards and charts were just pictures that you could not click or interact with.",
        "No Safe Transfer Method: Only had a basic 'Buy' button, with no answer to how code or money is actually protected."
    ], badge="ORIGINAL SITE", accent=ACCENT_AMBER)

    add_card(s5, 6.8, 1.9, 5.7, 5.0, "Our Innovation: Safe 4-Step Handover", [
        "The Safe Handover System: Solves the biggest fear in buying software using 4 clear steps:",
        "  1. Money Deposit: Buyer puts money in a safe lockbox (escrow).",
        "  2. Code Lock: Seller hands over code and cloud accounts into safe custody.",
        "  3. Domain Move: The website domain is moved to the buyer.",
        "  4. Testing & Release: Buyer tests the app; once confirmed, money goes to the seller.",
        "Interactive Simulator: Users can click through all 4 steps on the screen to see how it works.",
        "Clean, Professional Look: Simple typography and clean cards instead of confusing neon effects."
    ], badge="OUR INNOVATION", accent=ACCENT_GREEN)

    s5.notes_slide.notes_text_frame.text = (
        "Point 4: What is our main innovation compared to the original site? "
        "The original site had flashy neon colors, but the buttons and charts didn't do much, and there was no way to safely transfer apps. "
        "We built the Safe 4-Step Handover system: money is locked in escrow, code is secured, the domain is moved, "
        "and the buyer confirms everything works before the money is released. "
        "We also added a live simulator on every listing so users can click through and see the safety for themselves."
    )

    # -------------------------------------------------------------
    # SLIDE 6: Conclusion & Key Learnings
    # -------------------------------------------------------------
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6)
    add_header(s6, "5. Conclusion", "Conclusion & Key Learnings", "The three main things we learned while building this project.")

    add_card(s6, 0.8, 1.9, 3.6, 5.0, "1. Simple Code is Fast", [
        "We didn't need big frameworks like React to build an interactive multi-page website.",
        "Pure HTML, CSS, and JavaScript loaded in under 120ms with zero extra installs.",
        "Fewer libraries mean fewer bugs and zero maintenance headache."
    ], badge="CODE LESSON", accent=ACCENT_BLUE)

    add_card(s6, 4.86, 1.9, 3.6, 5.0, "2. Clear Info Beats Flashy UI", [
        "In finance and business, people care about honest numbers, not glowing neon buttons.",
        "Clean tables, verified Stripe data, and clear fee comparisons build real trust with buyers and sellers.",
        "Simple and plain design looks far more serious and professional."
    ], badge="DESIGN LESSON", accent=ACCENT_GREEN)

    add_card(s6, 8.93, 1.9, 3.6, 5.0, "3. Safety is What Sells", [
        "Anyone can build a website that lists apps for sale.",
        "The real challenge is making sure nobody gets scammed when handing over code and money.",
        "Building the safe 4-step handover is what makes this platform genuinely useful and valuable."
    ], badge="PRODUCT LESSON", accent=ACCENT_AMBER)

    s6.notes_slide.notes_text_frame.text = (
        "Point 5: To wrap up, our three biggest learnings: "
        "First, simple code is faster and easier to maintain than heavy frameworks. "
        "Second, clean numbers and honest data build much more trust than flashy animations. "
        "And third, solving the safety problem of handing over code is what gives this platform real value. "
        "The website is live and ready to view. Thank you!"
    )

    output_path = r"C:\Users\rohit\.gemini\antigravity\scratch\vventra-remake\vventra_presentation.pptx"
    prs.save(output_path)
    print(f"Presentation saved to: {output_path}")

if __name__ == "__main__":
    create_deck()
